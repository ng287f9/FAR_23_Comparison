#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
T6 — 按分部批量生成《14 CFR Part 23 条款修订历史与背景分析》Word 分卷。

与 §23.561 单条模板（build_23_561_doc.py）的分工：
  · 单条模板  = 精写版，含逐版叙事剖析与彩色条文 diff，人工投入大，只做重点条款
  · 本脚本    = 取证版，逐条款给出「年表 + 修订指令原文 + 讨论摘录 + 更正件提示
                + 参考文献 + 初始条文」，全部自动生成，覆盖分部全部条款

数据源（全部本地）：
  · FAR23_Amdt23-63_条款清单.csv       DRS 条款修订版记录
  · FAR23_条款-规则文件对应表.csv       (条款 × 修订案) → NPRM / Final Rule 编号
  · Rulemaking Docs/                   官方原文（part23_source.build_doc_index 索引）
  · FAR23_受影响条款清单.csv            更正件影响面（可选，用于提示）

用法：
    python build_subpart_doc.py                 # 默认 Subpart C
    python build_subpart_doc.py "Subpart B - Flight"
    python build_subpart_doc.py C --limit 5     # 只做前 5 条（调试）
"""
import argparse
import os
import re
import sys
from collections import OrderedDict, Counter

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

import docx_layout as L
import part23_source as S

ROOT = os.path.dirname(os.path.abspath(__file__))

SUBPART_KEY = {
    "A": "Subpart A - General",
    "B": "Subpart B - Flight",
    "C": "Subpart C - Structure",
    "D": "Subpart D - Design and Construction",
    "E": "Subpart E - Powerplant",
    "F": "Subpart F - Equipment",
    "G": "Subpart G - Operating Limitations and Information",
}
SUBPART_CN = {
    "A": "A 分部　总则", "B": "B 分部　飞行", "C": "C 分部　结构",
    "D": "D 分部　设计与构造", "E": "E 分部　动力装置", "F": "F 分部　设备",
    "G": "G 分部　使用限制与资料",
}


# ---------------------------------------------------------------- 更正件影响面
def load_correction_hits():
    """条款号 -> [更正件 FR 引证]"""
    import csv
    p = os.path.join(ROOT, "FAR23_受影响条款清单.csv")
    if not os.path.exists(p):
        return {}
    out = {}
    for r in csv.DictReader(open(p, encoding="utf-8-sig")):
        s = (r.get("条款/附录") or "").strip()
        frs = [x.strip() for x in (r.get("更正件FR引证") or "").split(";") if x.strip()]
        if s and frs:
            out[s] = frs
    return out


def _dedupe_key(blk):
    """同一段修订指令在 NPRM 与 Final Rule 里各出现一次（措辞几乎相同），
    用「去掉编号前缀 + 全部空白后的前 160 字符」做指纹。"""
    s = re.sub(r"\s+", "", "\n".join(blk))
    s = re.sub(r"^\d{1,3}[\.\)]", "", s)
    return s[:160]


def dedupe(items):
    """items: [(来源, 文件名, blk)] → 去重，同指纹保留 Final Rule 版。

    两轮：① 指纹完全相同；② 短块是长块的前缀（同一指令被截断成不同长度）。
    """
    seen = {}
    out = []
    for src, fn, blk in items:
        k = _dedupe_key(blk)
        if k in seen:
            i = seen[k]
            if out[i][0].startswith("NPRM") and src.startswith("Final"):
                out[i] = (src, fn, blk)
            continue
        seen[k] = len(out)
        out.append((src, fn, blk))

    keys = [_dedupe_key(b) for _, _, b in out]
    keep = []
    for i, k in enumerate(keys):
        # 若本块是某个更长块的前缀 → 丢弃本块
        drop = False
        for j, k2 in enumerate(keys):
            if i != j and len(k2) > len(k) and k2.startswith(k[:60]) and len(k) < 120:
                drop = True
                break
        if not drop:
            keep.append(out[i])
    return keep


# ---------------------------------------------------------------- 每条款取证
def collect_section(sec, meta, di, corr):
    """返回该条款的全部取证结果"""
    num = sec.replace("Sec. ", "").strip()
    out = {
        "num": num, "title": meta["title"], "versions": [],
        "init": None, "refs": OrderedDict(), "corrs": [],
    }
    for a, v in meta["versions"].items():
        nums = corr.get((sec, a), {})
        frs = nums.get("Final Rule", [])
        nps = nums.get("NPRM", [])
        A = S.resolve_docs(";".join(frs), "Final Rule", di, num) if frs else []
        N = S.resolve_docs(";".join(nps), "NPRM", di, num) if nps else []
        entry = {
            "amdt": a, "eff": v["eff"], "fr_actions": v["fr_actions"],
            "fr_docs": A, "nprm_docs": N,
            "amend": [], "discuss": [],
            "fr_nums": frs, "nprm_nums": nps,
        }
        for d in A:
            m = S.mine_doc(d, num)
            for b in m["amend"]:
                entry["amend"].append(("Final Rule", d.name, b))
            for b in m["discuss"][:3]:
                entry["discuss"].append(("Final Rule", d.name, b))
            out["refs"][d.name] = d
        for d in N:
            m = S.mine_doc(d, num)
            for b in m["amend"]:
                entry["amend"].append(("NPRM（提案）", d.name, b))
            for b in m["discuss"][:2]:
                entry["discuss"].append(("NPRM（提案）", d.name, b))
            out["refs"][d.name] = d
        entry["amend"] = dedupe(entry["amend"])
        entry["discuss"] = dedupe(entry["discuss"])
        out["versions"].append(entry)
    out["init"] = S.initial_clause_text(num)
    return out


# ---------------------------------------------------------------- 文档装配
def new_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.0)
    st = doc.styles["Normal"]
    st.font.name = L.EN_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(L.qn("w:eastAsia"), L.CN_FONT)
    st.paragraph_format.line_spacing = 1.35
    L.add_page_number_footer(sec)
    return doc


def cover(doc, sub_key, data):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    L.para_spacing(p, 40, 4)
    L.set_run(p.add_run("14 CFR Part 23 条款修订历史与背景分析"),
              size=22, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    L.para_spacing(p, 0, 6)
    L.set_run(p.add_run(SUBPART_CN.get(sub_key, sub_key)),
              size=17, bold=True, color=RGBColor(0x2E, 0x74, 0xB5))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    L.para_spacing(p, 0, 6)
    L.set_run(p.add_run(SUBPART_KEY.get(sub_key, sub_key)),
              size=11, color=RGBColor(0x59, 0x59, 0x59))

    nsec = len(data)
    nver = sum(len(d["versions"]) for d in data)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    L.para_spacing(p, 0, 18)
    L.set_run(p.add_run(f"快照：Part 23 @ Amdt 23-63 as of 08/29/2017 and earlier　｜　"
                        f"共 {nsec} 个条款 / {nver} 个历史版本"),
              size=11, color=RGBColor(0x59, 0x59, 0x59))

    L.info_box(doc, "文档说明与编制依据", [
        [("• 修订日期、修订案号、NPRM 与 Final Rule 编号：取自 ", 'normal'),
         ("FAA 动态法规系统（DRS）", 'bold'),
         ("「Part 23 @ Amdt 23-63 as of 08/29/2017 and earlier」条款修订版索引。", 'normal')],
        "• 修订指令原文与讨论/评论处置段落：逐条引自本地已下载的 NPRM 与 Final Rule 官方文本"
        "（Federal Register 影印件及提取文本），每条条款末尾列出所引文献。",
        "• 更正通告与 CFR 编纂校正：来源于官方 CFR 2016 版溯源注、federalregister.gov API 与"
        "govinfo 整期影印件三源交叉验证，见 FAR23_更正通告权威台账.csv。",
        "• 本卷为「取证版」——逐条款给出可核查的原始证据（年表 / 指令原文 / 讨论摘录 / 更正件提示），"
        "不做人工叙事发挥；重点条款另有精写版单条文档。",
    ])


def overview(doc, sub_key, data, corrhits):
    L.H(doc, "分部概览", 1)
    amdt = Counter()
    for d in data:
        for v in d["versions"]:
            amdt[v["amdt"]] += 1
    L.mk_table(doc, ["项目", "数值"], [
        ["分部", SUBPART_KEY.get(sub_key, sub_key)],
        ["条款数", str(len(data))],
        ["历史版本总数", str(sum(len(d["versions"]) for d in data))],
        ["平均每条款版本数", f"{sum(len(d['versions']) for d in data)/max(len(data),1):.2f}"],
        ["涉及的修订案数", str(len(amdt))],
        ["带更正件/校正的条款数",
         str(sum(1 for d in data if d["num"] in corrhits))],
        ["有修订指令原文的条款数",
         str(sum(1 for d in data if any(v["amend"] for v in d["versions"])))],
        ["有讨论/背景摘录的条款数",
         str(sum(1 for d in data if any(v["discuss"] for v in d["versions"])))],
        ["可回溯到 1965 年初版条文的条款数",
         str(sum(1 for d in data if d["init"]))],
    ], widths=[6.0, 10.8], font=9.5)

    L.P(doc, f"本分部涉及的修订案（按影响条款数排序，共 {len(amdt)} 个）：", size=10.5, bold=True)
    rows = [[a, str(n)] for a, n in sorted(amdt.items(), key=lambda kv: (-kv[1], kv[0]))]
    L.mk_table(doc, ["修订案", "影响条款数"], rows, widths=[4.0, 3.0], font=9)


def toc(doc, data):
    L.H(doc, "目录", 1)
    L.add_toc_field(doc)
    for d in data:
        p = doc.add_paragraph()
        L.para_spacing(p, 1, 1)
        L.set_run(p.add_run(f"§ {d['num']}　{d['title'][:60]}"),
                  size=10, color=RGBColor(0x1F, 0x4E, 0x79))
    doc.add_page_break()


def one_section(doc, d, corrhits):
    num, title = d["num"], d["title"]
    L.H(doc, f"§ {num}　{title}", 2)

    # ---- 基本信息
    nver = len(d["versions"])
    amdt_seq = " → ".join(v["amdt"] for v in d["versions"])
    L.mk_table(doc, ["项目", "内容"], [
        ["条款号", f"14 CFR § {num}"],
        ["标题", title],
        ["历史版本数", str(nver)],
        ["修订沿革", amdt_seq],
    ], widths=[3.2, 13.6], font=9)

    # ---- 修订年表
    L.H(doc, "修订年表", 3)
    rows = []
    for i, v in enumerate(d["versions"], 1):
        nprm = "; ".join(v["nprm_nums"]) or "—"
        fr = "; ".join(v["fr_nums"]) or "—"
        ev = []
        if v["amend"]:
            ev.append(f"修订指令 {len(v['amend'])} 段")
        if v["discuss"]:
            ev.append(f"讨论/评论 {len(v['discuss'])} 段")
        rows.append([str(i), v["amdt"], v["eff"] or "—", nprm, fr,
                     "；".join(ev) or "（本版为交叉引用所致，未改本条条文）"])
    L.mk_table(doc, ["#", "修订案", "生效日期", "NPRM 编号", "Final Rule 编号", "本地取证"],
               rows, widths=[0.9, 2.0, 2.2, 3.4, 3.4, 4.9], font=8.5,
               align_center_cols=(0, 1, 2))

    # ---- 更正件提示
    if num in corrhits:
        L.info_box(doc, "更正通告 / CFR 编纂校正提示", [
            "本条款的 CFR 溯源注或 FR 全库检索显示，除修订案之外还有以下更正类文献涉及本条：　"
            + "、".join(corrhits[num]),
            "⚠ 更正件通常改动拼写、交叉引用编号或表头，不产生新的 Amdt 号，"
            "DRS 条款版本索引与 CFR 溯源注都可能漏记；引用时请以现行 CFR 文本为准。",
        ], bg=L.C_SUB_BG)

    # ---- 修订指令原文
    has = [v for v in d["versions"] if v["amend"]]
    if has:
        L.H(doc, "修订指令原文（英文）", 3)
        for v in has:
            L.P(doc, f"Amdt {v['amdt']}（生效 {v['eff'] or '—'}）", size=10, bold=True,
                color=RGBColor(0x2E, 0x74, 0xB5))
            for src, fn, blk in v["amend"]:
                L.P(doc, f"［{src}］", size=8, italic=True,
                    color=RGBColor(0x80, 0x80, 0x80), indent=0.5, before=2, after=0)
                txt = "\n".join(blk)
                L.P(doc, txt[:2600] + ("…" if len(txt) > 2600 else ""),
                    size=8.5, indent=0.8)

    # ---- 讨论/背景摘录
    hasd = [v for v in d["versions"] if v["discuss"]]
    if hasd:
        L.H(doc, "修订背景与评论处置摘录（英文）", 3)
        for v in hasd:
            L.P(doc, f"Amdt {v['amdt']}", size=10, bold=True,
                color=RGBColor(0x2E, 0x74, 0xB5))
            for src, fn, blk in v["discuss"]:
                L.P(doc, f"［{src}］", size=8, italic=True,
                    color=RGBColor(0x80, 0x80, 0x80), indent=0.5, before=2, after=0)
                txt = "\n".join(blk)
                L.P(doc, txt[:2200] + ("…" if len(txt) > 2200 else ""),
                    size=8.5, indent=0.8)

    # ---- 初始版条文
    if d["init"]:
        L.H(doc, "1965 年初版条文（Docket 4080，英文）", 3)
        L.P(doc, "\n".join(d["init"])[:3000], size=8.5, indent=0.5)

    # ---- 参考文献
    if d["refs"]:
        L.H(doc, "本条引用文献", 3)
        rows = [[os.path.basename(n)] for n in d["refs"]]
        L.mk_table(doc, ["本地文件名"], rows, widths=[16.8], font=8)

    doc.add_paragraph()


def build(sub_key, limit=None, out=None):
    key = SUBPART_KEY.get(sub_key, sub_key)
    doc_index = S.build_doc_index()
    corr = S.load_correspondence()
    corrhits = load_correction_hits()
    secs = S.load_sections(key)

    data = []
    for i, (sec, meta) in enumerate(secs.items()):
        if limit and i >= limit:
            break
        data.append(collect_section(sec, meta, doc_index, corr))

    doc = new_doc()
    cover(doc, sub_key, data)
    toc(doc, data)
    overview(doc, sub_key, data, corrhits)
    for d in data:
        one_section(doc, d, corrhits)

    out = out or os.path.join(
        ROOT, f"FAR23_{SUBPART_KEY.get(sub_key, sub_key).split(' - ')[0].replace('Subpart ','Subpart')}"
              f"_修订史.docx")
    doc.save(out)
    print("已生成:", out)
    print(f"   条款 {len(data)}　段落 {len(doc.paragraphs)}　表格 {len(doc.tables)}")
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("subpart", nargs="?", default="C")
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    build(a.subpart, a.limit or None, a.out)
