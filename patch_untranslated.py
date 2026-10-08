#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把分部卷里「翻译失败退回英文」的段落补译并写回。

build_clause_doc.py 在 API 限流/额度耗尽时会把英文原文直接落到段落里。
本脚本扫描这些段落，取它们下方标注的英文原文，翻译后写回原位。

用法：
    python patch_untranslated.py <docx 路径> [--workers 2] [--rounds 3]

判定标准：段落以 "▪ 背景与评论处置" 开头且中文字符占比 < 15%。
英文原文取自该段之后 "英文原文：" 与下一个 "来源：" / "▪" / 标题之间的行。
"""
import re
import sys
import time
from pathlib import Path

from docx import Document

import translate as T

CJK = re.compile(r"[\u4e00-\u9fff]")
BULLET_PREFIX = "▪ 背景与评论处置"
EN_MARK = "英文原文："


def zh_ratio(s):
    return len(CJK.findall(s)) / max(len(s), 1)


def find_jobs(doc):
    """返回 [(段对象, 英文原文)]"""
    jobs = []
    cur, collecting, buf = None, False, []

    def flush():
        nonlocal cur, collecting, buf
        if cur is not None:
            en = "\n".join(x for x in buf if x.strip()).strip()
            if en:
                jobs.append((cur, en))
        cur, collecting, buf = None, False, []

    for pa in doc.paragraphs:
        t = pa.text.strip()
        if t.startswith(BULLET_PREFIX):
            flush()
            cur = pa if zh_ratio(t) < 0.15 else None
            continue
        if cur is None:
            continue
        if t.startswith(EN_MARK):
            collecting = True
            continue
        if not collecting:
            continue
        if (t.startswith("来源：") or t.startswith("▪")
                or t.startswith("修订指令英文原文")
                or pa.style.name.startswith("Heading")):
            flush()
            continue
        buf.append(t)
    flush()
    return jobs


def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--workers", type=int, default=2)
    ap.add_argument("--rounds", type=int, default=3)
    a = ap.parse_args()

    p = Path(a.docx)
    doc = Document(str(p))
    jobs = find_jobs(doc)
    if not jobs:
        print(f"{p.name}: 无需补译")
        return 0
    print(f"{p.name}: 待补译 {len(jobs)} 条", flush=True)

    tr = T.Translator(workers=a.workers, verbose=False)
    ok = 0
    for i, (para, en) in enumerate(jobs, 1):
        for _ in range(a.rounds):
            try:
                zh = tr.translate(en)
            except Exception:
                zh = None
            if zh and zh_ratio(zh) > 0.15:
                for r in para.runs[1:]:
                    r.text = ""
                if len(para.runs) > 1:
                    para.runs[1].text = zh
                else:
                    para.add_run(zh)
                ok += 1
                print(f"   [{i}/{len(jobs)}] OK", flush=True)
                break
            time.sleep(5)
        else:
            print(f"   [{i}/{len(jobs)}] 仍失败", flush=True)
        tr.save()
    doc.save(str(p))
    tr.save()
    print(f"{p.name}: 写回 {ok}/{len(jobs)}", flush=True)
    return 0 if ok == len(jobs) else 1


if __name__ == "__main__":
    sys.exit(main())
