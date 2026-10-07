#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Word 排版公共模块 —— 从 build_23_561_doc.py 抽出，供批量生成各分部卷复用。

着色约定：蓝=新增 / 红=修订 / 灰+删除线=删除 / 黑=继承未变
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement



# ---------------- 颜色 / 字体 ----------------
C_NEW = RGBColor(0x00, 0x70, 0xC0)      # 新增：蓝
C_REV = RGBColor(0xC0, 0x00, 0x00)      # 修订：红
C_DEL = RGBColor(0x80, 0x80, 0x80)      # 删除：灰
C_TXT = RGBColor(0x1A, 0x1A, 0x1A)      # 正文
C_HDR_BG = "1F4E79"                     # 表头底色
C_SUB_BG = "DDEBF7"                     # 交替/次级底色
C_NOTE_BG = "FFF2CC"                    # 提示框底色
C_BOX_BG = "F2F2F2"                     # 条文框底色
CN_FONT = "微软雅黑"
EN_FONT = "Calibri"


# ---------------- 底层工具 ----------------
def set_run(run, size=10.5, bold=False, color=C_TXT, italic=False,
            strike=False, cn=CN_FONT, en=EN_FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn)
    if strike:
        run.font.strike = True
    return run


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def cell_margins(table, top=60, bottom=60, left=90, right=90):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for tag, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + tag)
        e.set(qn('w:w'), str(val))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader')
    e.set(qn('w:val'), 'true')
    trPr.append(e)


def fill_runs(p, segs, size):
    """把 segs 渲染进段落 p。segs 可为 str、[(text, style)] 或二者的混合列表。"""
    if isinstance(segs, str):
        segs = [(segs, 'normal')]
    for item in segs:
        if isinstance(item, str):
            text, st = item, 'normal'
        else:
            text, st = item
        r = p.add_run(text)
        if st == 'new':
            set_run(r, size=size, bold=True, color=C_NEW)
        elif st == 'rev':
            set_run(r, size=size, bold=True, color=C_REV)
        elif st == 'del':
            set_run(r, size=size, color=C_DEL, strike=True)
        elif st == 'bold':
            set_run(r, size=size, bold=True, color=C_TXT)
        elif st == 'head':
            set_run(r, size=size, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
        else:
            set_run(r, size=size, color=C_TXT)
    return p


def para_spacing(p, before=4, after=4, line=1.35):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_toc_field(doc, levels="1-3"):
    """插入 TOC 域。levels 决定收录到第几级大纲（分部卷用 "1-2"，单条用 "1-3"）。"""
    p = doc.add_paragraph()
    r = p.add_run()
    for typ, txt in (('begin', None), ('separate', "　※ 在 Word/WPS 中按 Ctrl+A 后 F9（或右键→更新域）生成目录"),
                     ('end', None)):
        f = OxmlElement('w:fldChar')
        f.set(qn('w:fldCharType'), typ)
        r._r.append(f)
        if typ == 'begin':
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = f'TOC \\o "{levels}" \\h \\z \\u'
            r._r.append(it)
        elif typ == 'separate':
            t = OxmlElement('w:t')
            t.text = txt
            r._r.append(t)
    set_run(r, size=10, italic=True, color=C_DEL)
    return p


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p.add_run(FOOTER_TEXT + "　|　第 ")
    set_run(r1, size=8.5, color=C_DEL)
    r = p.add_run()
    f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), 'begin'); r._r.append(f)
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'; r._r.append(it)
    f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), 'end'); r._r.append(f)
    set_run(r, size=8.5, color=C_DEL)
    r2 = p.add_run(" 页")
    set_run(r2, size=8.5, color=C_DEL)


FOOTER_TEXT = "14 CFR Part 23 条款修订历史与背景分析"


# ---------------- 大纲级别 ----------------
# 标题既要有视觉层级，也必须是真正的 Word 大纲级别（Heading N 样式），
# 否则导航窗格与大纲视图是空的，长卷无法跳转。
HEADING_SPEC = {
    1: (16.0, RGBColor(0x1F, 0x4E, 0x79), 16, 8),
    2: (13.5, RGBColor(0x2E, 0x74, 0xB5), 12, 6),
    3: (12.0, RGBColor(0x40, 0x40, 0x40), 10, 5),
    4: (11.0, RGBColor(0x59, 0x59, 0x59), 8, 4),
}

_STYLED = set()  # 已配置过标题样式的文档 id，避免重复写


def setup_heading_styles(doc):
    """把内置 Heading 1–4 改成本项目配色（去下划线/斜体），并保持大纲级别。

    python-docx 默认模板的 Heading 样式是蓝色的 Theme 色 + 可能带下划线，
    这里统一改成中文字体 + 本项目配色，视觉上和手绘标题一致。
    """
    did = id(doc)
    if did in _STYLED:
        return
    _STYLED.add(did)
    for lv, (size, color, before, after) in HEADING_SPEC.items():
        st = doc.styles[f"Heading {lv}"]
        st.font.name = EN_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.underline = False
        st.font.color.rgb = color
        el = st.element
        rpr = el.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), CN_FONT)
        pf = st.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.keep_with_next = True
        pf.line_spacing = 1.2


def H(doc, text, level=1, size=None, color=None, before=14, after=6):
    """写一个带大纲级别的标题。level 1–4 → Word Heading 1–4。"""
    level = max(1, min(int(level), 4))
    setup_heading_styles(doc)
    p = doc.add_paragraph()
    p.style = doc.styles[f"Heading {level}"]
    dsize, dcolor, dbefore, dafter = HEADING_SPEC[level]
    pf = p.paragraph_format
    pf.space_before = Pt(before if before != 14 else dbefore)
    pf.space_after = Pt(after if after != 6 else dafter)
    pf.keep_with_next = True
    r = p.add_run(text)
    set_run(r, size=size or dsize, bold=True, color=color or dcolor)
    return p


def P(doc, text, size=10.5, indent=0, bold=False, color=C_TXT, before=3, after=3,
      italic=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    para_spacing(p, before, after)
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color, italic=italic)
    return p


def RICH(doc, segments, size=10.5, indent=0, before=3, after=3):
    """segments: [(text, style)]  style in normal/new/rev/del/bold"""
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    para_spacing(p, before, after)
    for text, st in segments:
        r = p.add_run(text)
        if st == 'new':
            set_run(r, size=size, bold=True, color=C_NEW)
        elif st == 'rev':
            set_run(r, size=size, bold=True, color=C_REV)
        elif st == 'del':
            set_run(r, size=size, color=C_DEL, strike=True)
        elif st == 'bold':
            set_run(r, size=size, bold=True, color=C_TXT)
        else:
            set_run(r, size=size, color=C_TXT)
    return p


def BULLET(doc, label, segments, indent=0.6, size=10.5):
    """label 为加粗小标题，segments 为 [(text, style)]"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    para_spacing(p, 3, 3)
    r = p.add_run("▪ " + label)
    set_run(r, size=size, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    for text, st in segments:
        rr = p.add_run(text)
        if st == 'new':
            set_run(rr, size=size, bold=True, color=C_NEW)
        elif st == 'rev':
            set_run(rr, size=size, bold=True, color=C_REV)
        elif st == 'del':
            set_run(rr, size=size, color=C_DEL, strike=True)
        else:
            set_run(rr, size=size, color=C_TXT)
    return p


def mk_table(doc, headers, rows, widths=None, font=9.5, header_bg=C_HDR_BG,
             zebra=True, align_center_cols=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    cell_margins(t)
    hdr = t.rows[0]
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_spacing(p, 2, 2)
        set_run(p.add_run(h), size=font, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, header_bg)
    repeat_header(hdr)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.text = ""
            p = c.paragraphs[0]
            para_spacing(p, 2, 2)
            if i in align_center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fill_runs(p, val, font)
            if zebra and ri % 2 == 1:
                shade(c, "F7F9FC")
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def info_box(doc, title, lines, bg=C_NOTE_BG):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    cell_margins(t)
    c = t.rows[0].cells[0]
    c.text = ""
    shade(c, bg)
    p = c.paragraphs[0]
    para_spacing(p, 2, 3)
    set_run(p.add_run(title), size=10.5, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
    for ln in lines:
        pp = c.add_paragraph()
        para_spacing(pp, 1, 1)
        if isinstance(ln, list):
            # [(text, style), ...]
            for text, st in ln:
                r = pp.add_run(text)
                if st == 'bold':
                    set_run(r, size=9.5, bold=True, color=C_TXT)
                elif st == 'new':
                    set_run(r, size=9.5, bold=True, color=C_NEW)
                elif st == 'rev':
                    set_run(r, size=9.5, bold=True, color=C_REV)
                elif st == 'del':
                    set_run(r, size=9.5, color=C_DEL, strike=True)
                else:
                    set_run(r, size=9.5, color=C_TXT)
        else:
            set_run(pp.add_run("　" + ln), size=9.5, color=C_TXT)
    doc.add_paragraph()
    return t


def legend(doc):
    p = doc.add_paragraph()
    para_spacing(p, 2, 6)
    for text, col, bold in (("■ 蓝色", C_NEW, True), ("＝本版新增　　", C_TXT, False),
                            ("■ 红色", C_REV, True), ("＝本版修订/替换　　", C_TXT, False),
                            ("■ 灰色删除线", C_DEL, True), ("＝本版删除　　", C_TXT, False),
                            ("■ 黑色", C_TXT, False), ("＝继承上版、未变化", C_TXT, False)):
        r = p.add_run(text)
        set_run(r, size=9.5, bold=bold, color=col, strike=(col == C_DEL and bold))
    return p


def clause_table(doc, title, rows, note=None):
    """rows: [(段落号, [(文本, style)], 状态标记)]"""
    P(doc, title, size=11, bold=True, color=RGBColor(0x1F, 0x4E, 0x79), before=10, after=3)
    if note:
        P(doc, note, size=9.5, italic=True, color=C_DEL, indent=0.3, before=0, after=4)
    t = doc.add_table(rows=1, cols=3)
    t.style = 'Table Grid'
    t.autofit = False
    cell_margins(t)
    hdr = t.rows[0]
    for i, h in enumerate(["段落", "条文内容（英文原文）", "状态"]):
        c = hdr.cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_spacing(p, 2, 2)
        set_run(p.add_run(h), size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(c, C_HDR_BG)
    repeat_header(hdr)
    for label, segs, status in rows:
        cells = t.add_row().cells
        # 段落号
        cells[0].text = ""
        p = cells[0].paragraphs[0]
        para_spacing(p, 2, 2)
        set_run(p.add_run(label), size=9.5, bold=True)
        shade(cells[0], C_BOX_BG)
        # 条文
        cells[1].text = ""
        p = cells[1].paragraphs[0]
        para_spacing(p, 2, 2)
        for text, st in segs:
            r = p.add_run(text)
            if st == 'new':
                set_run(r, size=9.5, bold=True, color=C_NEW)
            elif st == 'rev':
                set_run(r, size=9.5, bold=True, color=C_REV)
            elif st == 'del':
                set_run(r, size=9.5, color=C_DEL, strike=True)
            elif st == 'head':
                set_run(r, size=9.5, bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
            else:
                set_run(r, size=9.5, color=C_TXT)
        # 状态
        cells[2].text = ""
        p = cells[2].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para_spacing(p, 2, 2)
        col = {"新增": C_NEW, "修订": C_REV, "删除": C_DEL}.get(status, C_TXT)
        set_run(p.add_run(status), size=9, bold=(status != "—"), color=col)
        if status == "新增":
            shade(cells[2], "E7F1FA")
        elif status == "修订":
            shade(cells[2], "FBE9E9")
        elif status == "删除":
            shade(cells[2], "F0F0F0")
    for i, w in enumerate((1.6, 13.6, 1.6)):
        for row in t.rows:
            row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


# =====================================================================
