#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""分部卷合规自检：大纲级别、每条款六节、来源标注、漏译。

用法：python check_volume.py <docx>
输出首三行为单行摘要，供 watch_and_commit.sh 取用。
"""
import re
import sys
import collections

from docx import Document

CJK = re.compile(r"[\u4e00-\u9fff]")


def main(path):
    d = Document(path)
    txts = [p.text.strip() for p in d.paragraphs]
    cl = collections.OrderedDict()
    cur = None
    for x in d.paragraphs:
        t = x.text.strip()
        s = x.style.name
        if s == "Heading 1" and t.startswith("§ "):
            cur = t.split("　")[0]
            cl[cur] = 0
        elif s == "Heading 2" and cur:
            cl[cur] += 1
    lv = collections.Counter(x.style.name for x in d.paragraphs
                             if x.style.name.startswith("Heading"))
    ok6 = all(v == 6 for v in cl.values())
    disc = [t for t in txts if t.startswith("▪ 背景与评论处置")]
    bad = [t for t in disc if len(CJK.findall(t)) < len(t) * 0.15]
    src = len([t for t in txts if t.startswith("来源：")])

    print(f"条款 {len(cl)} 段落 {len(d.paragraphs)} 表格 {len(d.tables)}")
    print(f"大纲 H1={lv.get('Heading 1',0)} H2={lv.get('Heading 2',0)} "
          f"H3={lv.get('Heading 3',0)} 每条款六节 {ok6}")
    print(f"讨论段 {len(disc)} 来源标注 {src} 未译 {len(bad)}")
    if not ok6:
        print("!! 存在节数不为 6 的条款：",
              {k: v for k, v in cl.items() if v != 6})
    for t in bad[:5]:
        print("   未译：", t[:90])
    return 0 if (ok6 and not bad) else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1]))
